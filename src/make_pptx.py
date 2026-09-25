"""Render the same deck as make_slides.py, but to an editable REFERATAS.pptx.

Everything is emitted as native PowerPoint shapes — text frames and real tables —
so the deck can be edited normally. Layout coordinates are in points and match
the PDF version one to one (960 x 540 pt = 13.333 x 7.5 in).
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

from slides_content import build
from slides_lib import F, MARK, W, H, MARGIN, plain, register_fonts, wrap

OUT = Path(__file__).resolve().parent.parent / "REFERATAS.pptx"

INK = RGBColor(0x16, 0x20, 0x2B)
MUTED = RGBColor(0x5C, 0x6B, 0x7A)
ACCENT = RGBColor(0x1F, 0x4E, 0x79)
ACCENT_SOFT = RGBColor(0xE8, 0xEE, 0xF5)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FLAG = RGBColor(0xB0, 0x3A, 0x2E)
FLAG_SOFT = RGBColor(0xFB, 0xEB, 0xE9)
RULE = RGBColor(0xC9, 0xD4, 0xE0)

FACE = "Calibri"
MONO = "Consolas"

# Target body sizes. A slide whose content would not fit is shrunk by a factor
# found per slide, never below MIN_SCALE, so the deck stays readable but nothing
# runs off the bottom edge.
BODY = 14.0
SUB = 13.0
KV = 13.5
NOTE = 13.0
TABLE_WIDE = 12.0      # tables with up to 4 columns
TABLE_NARROW = 11.0    # tables with more columns
MIN_SCALE = 0.80
BOTTOM_LIMIT = H - 44  # keep clear of the footer


def emu(points):
    return Emu(int(points * 12700))


def y_emu(points_from_top):
    return emu(points_from_top)


def add_runs(para, text, size, color=INK, base_bold=False):
    """Split **bold** / *italic* / `code` markers into real PowerPoint runs."""
    for part in MARK.split(text):
        if not part:
            continue
        bold, italic, face = base_bold, False, FACE
        if part.startswith("**"):
            bold, part = True, part[2:-2]
        elif part.startswith("`"):
            face, part = MONO, part[1:-1]
        elif part.startswith("*"):
            italic, part = True, part[1:-1]
        run = para.add_run()
        run.text = part
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.name = face
        run.font.color.rgb = color


def textbox(slide, x, top, width, lines, size, color=INK, space_after=4,
            height=None):
    """lines: list of (text, indent_level, bullet_char or None).

    `height` should be the measured height of the wrapped text; giving the shape
    its true height keeps the box outline tight when the deck is edited.
    """
    if height is None:
        height = sum(measure(t, size, width) * (size + space_after) for t, _, _ in lines)
        height = max(height, size + space_after)
    box = slide.shapes.add_textbox(emu(x), y_emu(top), emu(width), emu(height))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    first = True
    for text, indent, bullet in lines:
        para = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        para.space_after = Pt(space_after)
        para.line_spacing = 1.02
        if indent:
            para.left_indent = emu(indent + 14)
            para.first_line_indent = emu(-14)
        add_runs(para, (f"{bullet}  " if bullet else "") + text, size, color)
    return box


def measure(text, size, width, font=None):
    """Line count for a wrapped string, using the same metrics as the PDF."""
    return len(wrap(plain(text), font or F["regular"], size, width))


class PptxDeck:
    """Same interface as slides_lib.Deck, but writes a .pptx."""

    def __init__(self):
        register_fonts()
        self.prs = Presentation()
        self.prs.slide_width = emu(W)
        self.prs.slide_height = emu(H)
        self.blank = self.prs.slide_layouts[6]
        self.n = 0
        self.shrunk = []       # slides that had to be scaled down to fit

    # ---------------------------------------------------------------- chrome
    def _new(self):
        self.n += 1
        return self.prs.slides.add_slide(self.blank)

    def _rule(self, slide, top):
        bar = slide.shapes.add_shape(1, emu(MARGIN), y_emu(top),
                                     emu(W - 2 * MARGIN), emu(0.9))
        bar.fill.solid()
        bar.fill.fore_color.rgb = RULE
        bar.line.fill.background()
        bar.shadow.inherit = False

    def _header(self, slide, title, kicker):
        top = 44.0
        if kicker:
            textbox(slide, MARGIN, top, W - 2 * MARGIN,
                    [(kicker.upper(), 0, None)], 10, ACCENT)
            top += 20
        n_lines = measure(title, 21, W - 2 * MARGIN, F["bold"])
        box = textbox(slide, MARGIN, top, W - 2 * MARGIN, [], 21)
        para = box.text_frame.paragraphs[0]
        add_runs(para, title, 21, INK, base_bold=True)
        top += n_lines * 27 + 8
        self._rule(slide, top)
        return top + 22

    def _footer(self, slide, tag):
        if tag:
            textbox(slide, MARGIN, H - 32, W - 2 * MARGIN - 40,
                    [(tag, 0, None)], 8, MUTED)
        box = textbox(slide, W - MARGIN - 40, H - 32, 40,
                      [(str(self.n), 0, None)], 8, MUTED)
        box.text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT

    # ---------------------------------------------------------------- slides
    def title_slide(self, title, subtitle, meta):
        slide = self._new()
        bar = slide.shapes.add_shape(1, 0, 0, emu(14), emu(H))
        bar.fill.solid()
        bar.fill.fore_color.rgb = ACCENT
        bar.line.fill.background()
        bar.shadow.inherit = False

        x, width = MARGIN + 26, W - 2 * MARGIN - 40
        top = 150.0
        box = textbox(slide, x, top, width, [], 34)
        add_runs(box.text_frame.paragraphs[0], title, 34, INK, base_bold=True)
        top += measure(title, 34, width, F["bold"]) * 44 + 14

        textbox(slide, x, top, width, [(subtitle, 0, None)], 16, MUTED)
        top += measure(subtitle, 16, width) * 23 + 28

        self._rule(slide, top)
        top += 20
        textbox(slide, x, top, width, [(m, 0, None) for m in meta], 13, INK,
                space_after=6)

    def _run(self, slide, blocks, y, scale):
        """Lay out the body. With slide=None nothing is drawn, only measured."""
        pending = []           # consecutive text lines share one text box

        def flush(y):
            if not pending:
                return y
            height = sum(h for _, h in pending)
            if slide is not None:
                textbox(slide, MARGIN, y, W - 2 * MARGIN,
                        [ln for ln, _ in pending], BODY * scale, height=height)
            pending.clear()
            return y + height

        for block in blocks:
            kind = block[0]
            if kind == "b":
                size = (block[2] if len(block) > 2 else BODY) * scale
                h = measure(block[1], size, W - 2 * MARGIN - 16) * (size + 4.5) + 5
                pending.append(((block[1], 0, "•"), h))
            elif kind == "sub":
                size = SUB * scale
                h = measure(block[1], size, W - 2 * MARGIN - 40) * (size + 4.5) + 4
                pending.append(((block[1], 22, "–"), h))
            elif kind == "p":
                size = BODY * scale
                h = measure(block[1], size, W - 2 * MARGIN) * (size + 5) + 6
                pending.append(((block[1], 0, None), h))
            elif kind == "kv":
                y = flush(y)
                y = self._kv(slide, block[1], y, scale)
            elif kind == "table":
                y = flush(y)
                widths = block[3] if len(block) > 3 else None
                y = self._table(slide, block[1], block[2], y, widths, scale)
            elif kind == "note":
                y = flush(y)
                y = self._note(slide, block[1], block[2], y, scale)
            elif kind == "gap":
                y = flush(y) + block[1] * scale
        return flush(y)

    def slide(self, title, blocks, kicker=None, tag=None):
        slide = self._new()
        top = self._header(slide, title, kicker)

        scale = 1.0
        while scale > MIN_SCALE:
            if self._run(None, blocks, top, scale) <= BOTTOM_LIMIT:
                break
            scale -= 0.02
        scale = max(scale, MIN_SCALE)

        self._run(slide, blocks, top, scale)
        self._footer(slide, tag)
        if scale < 0.999:
            self.shrunk.append((self.n, round(scale, 2)))

    # ---------------------------------------------------------------- blocks
    def _kv(self, slide, rows, y, scale=1.0):
        size = KV * scale
        key_w = 168 * scale + 20
        line = size + 4.5
        for key, val in rows:
            val_w = W - 2 * MARGIN - key_w
            h = measure(val, size, val_w) * line
            if slide is not None:
                box = textbox(slide, MARGIN, y, key_w - 8, [], size,
                              height=max(h, line))
                add_runs(box.text_frame.paragraphs[0], key, size, ACCENT,
                         base_bold=True)
                textbox(slide, MARGIN + key_w, y, val_w, [(val, 0, None)], size,
                        height=h)
            y += h + 4
        return y + 4

    def _table(self, slide, header, rows, y, widths=None, scale=1.0):
        avail = W - 2 * MARGIN
        n = len(header)
        widths = widths or [avail / n] * n
        norm = avail / sum(widths)
        widths = [w * norm for w in widths]
        size = (TABLE_WIDE if n <= 4 else TABLE_NARROW) * scale

        def row_height(cells, font):
            lines = max(max(measure(str(c).split("\n")[k], size, widths[i] - 12, font)
                            for k in range(len(str(c).split("\n"))))
                        + str(c).count("\n") for i, c in enumerate(cells))
            return lines * (size + 3.4) + 9

        heights = [row_height(header, F["bold"])] + \
                  [row_height(r, F["regular"]) for r in rows]
        total = sum(heights)
        if slide is None:
            return y + total + 12

        shape = slide.shapes.add_table(len(rows) + 1, n, emu(MARGIN), y_emu(y),
                                       emu(avail), emu(total))
        table = shape.table
        table.first_row = False
        table.horz_banding = False
        for i, w in enumerate(widths):
            table.columns[i].width = emu(w)
        for r, h in enumerate(heights):
            table.rows[r].height = emu(h)

        def fill_cell(cell, text, bold, bg):
            cell.fill.solid()
            cell.fill.fore_color.rgb = bg
            cell.margin_left = cell.margin_right = emu(6)
            cell.margin_top = cell.margin_bottom = emu(3)
            cell.vertical_anchor = MSO_ANCHOR.TOP
            tf = cell.text_frame
            tf.word_wrap = True
            colour = WHITE if bold else INK
            for k, chunk in enumerate(str(text).split("\n")):
                para = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
                para.space_after = Pt(0)
                para.line_spacing = 1.0
                add_runs(para, chunk, size, colour, base_bold=bold)

        for i, cell_text in enumerate(header):
            fill_cell(table.cell(0, i), cell_text, True, ACCENT)
        for r, row in enumerate(rows, start=1):
            bg = ACCENT_SOFT if r % 2 else WHITE
            for i in range(n):
                fill_cell(table.cell(r, i), row[i] if i < len(row) else "", False, bg)

        return y + total + 12

    def _note(self, slide, label, text, y, scale=1.0):
        size = NOTE * scale
        avail = W - 2 * MARGIN - 28
        height = measure(text, size, avail) * (size + 4) + 30
        if slide is None:
            return y + height + 12

        panel = slide.shapes.add_shape(1, emu(MARGIN), y_emu(y),
                                       emu(W - 2 * MARGIN), emu(height))
        panel.fill.solid()
        panel.fill.fore_color.rgb = FLAG_SOFT
        panel.line.fill.background()
        panel.shadow.inherit = False

        edge = slide.shapes.add_shape(1, emu(MARGIN), y_emu(y), emu(3.5), emu(height))
        edge.fill.solid()
        edge.fill.fore_color.rgb = FLAG
        edge.line.fill.background()
        edge.shadow.inherit = False

        box = textbox(slide, MARGIN + 14, y + 7, avail, [], 10 * scale)
        add_runs(box.text_frame.paragraphs[0], label.upper(), 10 * scale, FLAG,
                 base_bold=True)
        textbox(slide, MARGIN + 14, y + 23, avail, [(text, 0, None)], size)
        return y + height + 12

    def save(self):
        self.prs.save(str(OUT))


def main():
    d = PptxDeck()
    build(d)
    d.save()
    print(f"wrote {OUT}  ({d.n} slides)")
    print(f"body size: {BODY} pt (tables {TABLE_WIDE}/{TABLE_NARROW} pt)")
    if d.shrunk:
        print("scaled down to fit:")
        for n, s in d.shrunk:
            print(f"  slide {n:>2}  x{s}  -> body {BODY * s:.1f} pt")
    else:
        print("no slide needed scaling")


if __name__ == "__main__":
    main()
