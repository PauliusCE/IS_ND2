"""Convert ATASKAITA.md to ATASKAITA.docx.

A small, purpose-built Markdown converter: this report uses headings, tables,
fenced code blocks, bullet and numbered lists, and inline bold/code. Pandoc is
not installed on this machine, so the subset is handled directly.
"""

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "ATASKAITA.md"
OUT = ROOT / "ATASKAITA.docx"

HEADER_FILL = "D9E2F3"
CODE_FILL = "F2F2F2"


def shade(cell_or_para, fill):
    """Apply background shading to a table cell or a paragraph."""
    el = cell_or_para._tc if hasattr(cell_or_para, "_tc") else cell_or_para._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    (el.get_or_add_tcPr() if hasattr(cell_or_para, "_tc") else el).append(shd)


TOKEN = re.compile(r"(\*\*.+?\*\*|\*[^*]+\*|`[^`]+`|\[[^\]]+\]\([^)]+\))", re.S)


def add_runs(paragraph, text, bold=False, italic=False):
    """Render inline **bold**, *italic*, `code` and [label](url).

    Recurses so that formatting nests correctly, e.g. **bold with `code`**.
    """
    for part in TOKEN.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            add_runs(paragraph, part[2:-2], bold=True, italic=italic)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            add_runs(paragraph, part[1:-1], bold=bold, italic=True)
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(0xA3, 0x1D, 0x1D)
            run.bold, run.italic = bold, italic
        elif part.startswith("["):
            run = paragraph.add_run(part[1:part.index("]")])
            run.underline = True
            run.bold, run.italic = bold, italic
        else:
            run = paragraph.add_run(part.replace("\\|", "|"))
            run.bold, run.italic = bold, italic


def split_row(line):
    return [c.strip() for c in re.split(r"(?<!\\)\|", line.strip())[1:-1]]


def main():
    lines = SRC.read_text(encoding="utf-8").splitlines()
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    style.paragraph_format.space_after = Pt(6)

    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # ---- fenced code block ----
        if stripped.startswith("```"):
            i += 1
            block = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            i += 1
            for text in block:
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.left_indent = Pt(12)
                run = p.add_run(text if text.strip() else " ")
                run.font.name = "Consolas"
                run.font.size = Pt(9)
                shade(p, CODE_FILL)
            doc.add_paragraph().paragraph_format.space_after = Pt(0)
            continue

        # ---- table ----
        if stripped.startswith("|") and i + 1 < len(lines) and \
                re.match(r"^\|[\s:\-|]+\|$", lines[i + 1].strip()):
            header = split_row(stripped)
            i += 2
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(split_row(lines[i].strip()))
                i += 1

            table = doc.add_table(rows=1, cols=len(header))
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for c, text in enumerate(header):
                cell = table.rows[0].cells[c]
                cell.paragraphs[0].text = ""
                add_runs(cell.paragraphs[0], text)
                for run in cell.paragraphs[0].runs:
                    run.bold = True
                    run.font.size = Pt(9)
                shade(cell, HEADER_FILL)
            for row in rows:
                cells = table.add_row().cells
                for c, text in enumerate(row[:len(header)]):
                    cells[c].paragraphs[0].text = ""
                    add_runs(cells[c].paragraphs[0], text)
                    for run in cells[c].paragraphs[0].runs:
                        run.font.size = Pt(9)
            doc.add_paragraph()
            continue

        # ---- headings ----
        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            text = stripped[level:].strip()
            if level == 1:
                h = doc.add_heading(text, 0)
                h.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                doc.add_heading(text, min(level - 1, 4))
            i += 1
            continue

        # ---- horizontal rule ----
        if stripped in ("---", "***", "___"):
            p = doc.add_paragraph()
            pPr = p._p.get_or_add_pPr()
            borders = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "6")
            bottom.set(qn("w:color"), "BFBFBF")
            borders.append(bottom)
            pPr.append(borders)
            i += 1
            continue

        # ---- lists ----
        m = re.match(r"^(\s*)[-*]\s+(.*)$", line)
        if m:
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.left_indent = Pt(18 + 18 * (len(m.group(1)) // 2))
            add_runs(p, m.group(2))
            i += 1
            continue

        m = re.match(r"^(\s*)\d+\.\s+(.*)$", line)
        if m:
            p = doc.add_paragraph(style="List Number")
            add_runs(p, m.group(2))
            i += 1
            continue

        # ---- blank / paragraph ----
        if not stripped:
            i += 1
            continue

        # join wrapped lines into one paragraph
        para = [stripped]
        i += 1
        while i < len(lines):
            nxt = lines[i].strip()
            if not nxt or nxt.startswith(("#", "|", "-", "*", "```", "---")) or \
                    re.match(r"^\d+\.\s", nxt):
                break
            para.append(nxt)
            i += 1
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_runs(p, " ".join(para))

    doc.save(OUT)
    print(f"wrote {OUT}")
    print(f"  paragraphs: {len(doc.paragraphs)}, tables: {len(doc.tables)}")


if __name__ == "__main__":
    sys.exit(main())
